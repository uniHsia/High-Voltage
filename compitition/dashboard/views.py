from django.shortcuts import render, redirect
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
import io
import logging
from functools import wraps
from pathlib import Path
from datetime import datetime
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Lazy import torch to avoid import errors if torch is not installed
try:
    import torch
except ImportError:
    torch = None

from .utils import (
    user_exists,
    register_user,
    verify_user,
)

# Lazy import model functions to avoid startup issues
def get_diagnosis_data():
    from .main_Diagnosis_1 import get_diagnosis_data as _func
    return _func()

def get_health_prediction_data():
    from .main_predict_2 import get_health_prediction_data as _func
    return _func()

logger = logging.getLogger(__name__)


def login_required_custom(view_func):
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.session.get('is_logged_in'):
            if request.path.startswith('/api/'):
                return JsonResponse(
                    {
                        'error': '未登录或会话已过期，请重新登录。',
                        'redirect_url': '/',
                    },
                    status=401,
                )

            messages.warning(request, '请先登录后再访问系统。')
            return redirect('login')
        return view_func(request, *args, **kwargs)
    return wrapper


def login_view(request):
    if request.session.get('is_logged_in'):
        return redirect('home')

    context = {
        'saved_username': '',
    }

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '').strip()

        context['saved_username'] = username

        if not username or not password:
            messages.error(request, '请完整填写用户名和密码。')
            return render(request, 'dashboard/login.html', context)

        if not user_exists(username):
            messages.error(request, '该用户尚未注册，请先注册。')
            return redirect('register')

        if verify_user(username, password):
            request.session['is_logged_in'] = True
            request.session['username'] = username
            messages.success(request, '登录成功，欢迎进入系统。')
            return redirect('home')
        else:
            messages.error(request, '用户名或密码错误。')
            return render(request, 'dashboard/login.html', context)

    return render(request, 'dashboard/login.html', context)


def register_view(request):
    context = {
        'saved_username': '',
        'saved_email': '',
    }

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        email = request.POST.get('email', '').strip()
        password = request.POST.get('password', '').strip()
        confirm_password = request.POST.get('confirm_password', '').strip()

        context['saved_username'] = username
        context['saved_email'] = email

        if not username or not email or not password or not confirm_password:
            messages.error(request, '请完整填写注册信息。')
            return render(request, 'dashboard/register.html', context)

        if password != confirm_password:
            messages.error(request, '两次输入的密码不一致。')
            return render(request, 'dashboard/register.html', context)

        if user_exists(username):
            messages.error(request, '该用户名已存在，请更换用户名。')
            return render(request, 'dashboard/register.html', context)

        register_user(username, email, password)
        messages.success(request, '注册成功，请登录。')
        return redirect('login')

    return render(request, 'dashboard/register.html', context)


def logout_view(request):
    request.session.flush()
    messages.success(request, '您已退出登录。')
    return redirect('login')


def select_device(request):
    """选择设备，更新会话中的选中设备"""
    if request.method == 'POST':
        device_name = request.POST.get('device_name', 'GIS-01')
        request.session['selected_device'] = device_name
        return JsonResponse({'success': True, 'device': device_name})
    return JsonResponse({'error': '仅支持POST请求'}, status=405)


def get_device_data(request):
    """获取选中设备的数据"""
    if request.method == 'GET':
        selected_device = request.session.get('selected_device', 'GIS-01')
        # 根据设备返回对应的数据
        device_data = {
            'GIS-01': {'name': 'GIS-01', 'score': 92.5, 'status': '正常'},
            'GIS-02': {'name': 'GIS-02', 'score': 67.2, 'status': '注意'},
            '变压器-01': {'name': '变压器-01', 'score': 88.3, 'status': '正常'},
            '电缆-01': {'name': '电缆-01', 'score': 94.1, 'status': '正常'},
        }
        data = device_data.get(selected_device, device_data['GIS-01'])
        return JsonResponse(data)
    return JsonResponse({'error': '仅支持GET请求'}, status=405)


@login_required_custom
def home(request):
    context = {
        'page_title': '首页',
        'health_score': 92.5,
        'health_status': '良好',
        'current_alarm_count': 3,
        'today_alarm_count': 12,
        'online_devices': '8/10',
        'device_list': [
            {'name': 'GIS-01', 'score': 92.5, 'status': '正常', 'checked': True},
            {'name': 'GIS-02', 'score': 67.2, 'status': '注意', 'checked': False},
            {'name': '变压器-01', 'score': 88.3, 'status': '正常', 'checked': False},
            {'name': '电缆-01', 'score': 94.1, 'status': '正常', 'checked': False},
        ],
        'latest_wave_records': [
            '[2025-06-01 14:23]',
            '[2025-06-01 13:15]',
        ],
        'latest_alarms': [
            '[2025-06-01 14:23] GIS-02 局放幅值超阈值 > 80dB（预警）',
            '[2025-06-01 13:15] 变压器-01 油温异常上升（注意）',
        ]
    }
    return render(request, 'dashboard/home.html', context)


@login_required_custom
def diagnosis(request):
    context = {
        'page_title': '智能诊断',
        'discharge_type': '金属颗粒',
        'confidence': 94.3,
        'severity': '中等',
        'feature_importance': [
            {'name': '相位区间特征贡献度', 'value': 78},
            {'name': '幅值特征', 'value': 25},
            {'name': '脉冲重复率', 'value': 41},
        ],
        'similar_cases': [
            '[2025-06-01 14:23]',
            '[2025-06-01 13:15]',
        ]
    }
    return render(request, 'dashboard/diagnosis.html', context)


@login_required_custom
def health(request):
    context = {
        'page_title': '健康评估',
        'device_name': 'GIS-01',
        'hi_score': 92.5,
        'status': '良好',
        'threshold': 80,
        'rul_days': 186,
        'rul_months': 6.2,
        'model_name': 'Informer',
        'predict_time': '2025-06-01',
    }
    return render(request, 'dashboard/health.html', context)


@login_required_custom
def history(request):
    context = {
        'page_title': '历史追溯',
        'selected_device': 'GIS-01',
        'data_type': '全部',
        'start_date': '2025-05-01',
        'end_date': '2025-06-01',
        'records': [
            {'time': '2025-05-28 10:23', 'type': '悬浮放电', 'confidence': '94.2%', 'hi': '86.6%', 'action': '查看'},
            {'time': '2025-05-30 15:10', 'type': '金属颗粒', 'confidence': '84.2%', 'hi': '81.5%', 'action': '查看'},
        ],
        'report_preview_text': None,
    }
    return render(request, 'dashboard/history.html', context)

def get_history_report_data():
    return {
        'selected_device': 'GIS-01',
        'data_type': '全部',
        'start_date': '2025-05-01',
        'end_date': '2025-06-01',
        'records': [
            {'time': '2025-05-28 10:23', 'type': '悬浮放电', 'confidence': '94.2%', 'hi': '86.6%', 'action': '查看'},
            {'time': '2025-05-30 15:10', 'type': '金属颗粒', 'confidence': '84.2%', 'hi': '81.5%', 'action': '查看'},
        ]
    }


@login_required_custom
def report_preview(request):
    report_data = get_history_report_data()

    preview_text = []
    preview_text.append("高电压设备智能健康管理一体化平台 - 历史诊断报告")
    preview_text.append(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    preview_text.append(f"设备名称：{report_data['selected_device']}")
    preview_text.append(f"数据类型：{report_data['data_type']}")
    preview_text.append(f"时间范围：{report_data['start_date']} 至 {report_data['end_date']}")
    preview_text.append("")
    preview_text.append("历史诊断记录：")

    for idx, item in enumerate(report_data['records'], start=1):
        preview_text.append(
            f"{idx}. 时间：{item['time']}，放电类型：{item['type']}，"
            f"置信度：{item['confidence']}，HI：{item['hi']}"
        )

    preview_result = "\n".join(preview_text)

    context = {
        'page_title': '历史追溯',
        'selected_device': report_data['selected_device'],
        'data_type': report_data['data_type'],
        'start_date': report_data['start_date'],
        'end_date': report_data['end_date'],
        'records': report_data['records'],
        'report_preview_text': preview_result,
    }
    return render(request, 'dashboard/history.html', context)


@login_required_custom
def report_export(request):
    report_format = request.GET.get('format', 'pdf').lower()
    report_data = get_history_report_data()

    if report_format == 'docx':
        return export_report_docx(report_data)
    return export_report_pdf(report_data)


def export_report_docx(report_data):
    document = Document()
    document.add_heading('高电压设备智能健康管理一体化平台 - 历史诊断报告', 0)

    document.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    document.add_paragraph(f"设备名称：{report_data['selected_device']}")
    document.add_paragraph(f"数据类型：{report_data['data_type']}")
    document.add_paragraph(f"时间范围：{report_data['start_date']} 至 {report_data['end_date']}")

    document.add_heading('历史诊断记录', level=1)

    table = document.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '时间'
    hdr_cells[1].text = '放电类型'
    hdr_cells[2].text = '置信度'
    hdr_cells[3].text = 'HI'

    for item in report_data['records']:
        row_cells = table.add_row().cells
        row_cells[0].text = item['time']
        row_cells[1].text = item['type']
        row_cells[2].text = item['confidence']
        row_cells[3].text = item['hi']

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = 'attachment; filename="history_report.docx"'
    return response


def export_report_pdf(report_data):
    buffer = io.BytesIO()

    # Windows 常见中文字体路径，优先尝试微软雅黑
    font_candidates = [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\msyh.ttf",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
        r"C:\Windows\Fonts\simsun.ttf",
    ]

    font_registered = False
    for font_path in font_candidates:
        try:
            pdfmetrics.registerFont(TTFont("CN_FONT", font_path))
            font_registered = True
            break
        except Exception:
            continue

    p = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    y = height - 50

    if font_registered:
        title_font = "CN_FONT"
        body_font = "CN_FONT"
    else:
        title_font = "Helvetica-Bold"
        body_font = "Helvetica"

    p.setFont(title_font, 14)
    p.drawString(50, y, "高电压设备智能健康管理一体化平台 - 历史诊断报告")

    y -= 30
    p.setFont(body_font, 11)
    p.drawString(50, y, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    y -= 20
    p.drawString(50, y, f"设备名称：{report_data['selected_device']}")
    y -= 20
    p.drawString(50, y, f"数据类型：{report_data['data_type']}")
    y -= 20
    p.drawString(50, y, f"时间范围：{report_data['start_date']} 至 {report_data['end_date']}")

    y -= 35
    p.setFont(title_font, 12)
    p.drawString(50, y, "历史诊断记录")

    y -= 25
    p.setFont(body_font, 10)

    for idx, item in enumerate(report_data['records'], start=1):
        line = (
            f"{idx}. 时间：{item['time']}  放电类型：{item['type']}  "
            f"置信度：{item['confidence']}  HI：{item['hi']}"
        )
        p.drawString(50, y, line[:95])  # 简单防止一行过长
        y -= 18

        if y < 60:
            p.showPage()
            y = height - 50
            p.setFont(body_font, 10)

    p.save()
    buffer.seek(0)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="history_report.pdf"'
    return response


@login_required_custom
def run_diagnosis_model(request):
    """运行智能诊断模型并返回JSON数据"""
    if request.method not in ('GET', 'POST'):
        return JsonResponse({'error': '仅支持 GET/POST 请求'}, status=405)

    try:
        # 调用模型获取数据
        logger.info("开始运行诊断模型...")
        data = get_diagnosis_data()
        logger.info(f"诊断模型运行成功，结果：{data.get('discharge_type', 'unknown')}")
        return JsonResponse(data, safe=False)
    except FileNotFoundError as e:
        logger.error(f"模型文件未找到：{str(e)}")
        return JsonResponse({'error': '模型文件未找到，请确保模型已正确训练和保存'}, status=500)
    except torch.cuda.OutOfMemoryError:
        logger.error("GPU内存不足")
        return JsonResponse({'error': 'GPU内存不足，请尝试关闭其他程序或使用CPU模式'}, status=500)
    except Exception as e:
        logger.exception(f"诊断模型运行失败：{str(e)}")
        return JsonResponse({'error': f'诊断模型运行失败：{str(e)}'}, status=500)


@login_required_custom
def run_health_model(request):
    """运行健康评估模型并返回JSON数据"""
    if request.method not in ('GET', 'POST'):
        return JsonResponse({'error': '仅支持 GET/POST 请求'}, status=405)

    try:
        # 检查是否已有训练好的模型
        model_path = Path(__file__).resolve().parent.parent / 'improved_cross_attention_health_assessment.pth'

        if model_path.exists():
            logger.info("发现已训练的模型，加载中...")
        else:
            logger.info("未发现已训练模型，开始训练（这可能需要1-2分钟）...")

        # 调用模型获取数据
        logger.info("开始运行健康评估模型...")
        data = get_health_prediction_data()

        logger.info(f"健康评估模型运行成功，HI指数: {data['hi_score']}")
        return JsonResponse(data, safe=False)

    except FileNotFoundError as e:
        logger.error(f"模型文件未找到：{str(e)}")
        return JsonResponse({'error': '模型文件未找到，请确保模型已正确训练和保存'}, status=500)
    except torch.cuda.OutOfMemoryError:
        logger.error("GPU内存不足")
        return JsonResponse({'error': 'GPU内存不足，请尝试关闭其他程序或使用CPU模式'}, status=500)
    except Exception as e:
        import traceback
        logger.error(f"健康评估模型运行失败: {str(e)}\n{traceback.format_exc()}")
        return JsonResponse({'error': f'健康评估模型运行失败：{str(e)}'}, status=500)


def test_import(request):
    """测试模型是否能成功导入"""
    try:
        from .main_Diagnosis_1 import get_diagnosis_data as test1
        from .main_predict_2 import get_health_prediction_data as test2
        return JsonResponse({
            'success': True,
            'message': '所有模型函数已成功导入'
        })
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e) + '\n' + traceback.format_exc()
        })


def test_page(request):
    """测试页面"""
    return render(request, 'dashboard/test_api.html')
